from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from supabase import create_client
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from docx import Document
from docx.shared import Pt, RGBColor
from dotenv import load_dotenv
from dateutil import parser as dateutil_parser
import os
import tempfile

load_dotenv()

router = APIRouter()

_supabase = None

def get_supabase():
    global _supabase
    if _supabase is None:
        _supabase = create_client(os.getenv("SUPABASE_URL"), os.getenv("SUPABASE_SERVICE_KEY"))
    return _supabase


def fetch_export_data(experiment_id: str, include_transcripts: bool):
    sb = get_supabase()

    exp = sb.table("experiments").select("*").eq("id", experiment_id).single().execute()
    if not exp.data:
        raise HTTPException(status_code=404, detail="Experiment not found")

    subs = sb.table("sub_experiments").select("*").eq("experiment_id", experiment_id).execute()

    sub_ids = [s["id"] for s in subs.data]
    entries = []
    if sub_ids:
        entries = sb.table("data_entries").select("*").in_("sub_experiment_id", sub_ids).order("created_at").execute().data

    return exp.data, subs.data, entries


def build_pdf(experiment_id: str, include_transcripts: bool) -> str:
    exp, subs, entries = fetch_export_data(experiment_id, include_transcripts)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    doc = SimpleDocTemplate(tmp.name, pagesize=letter, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle("title", parent=styles["Heading1"], fontSize=20, spaceAfter=4)
    story.append(Paragraph(exp["title"], title_style))

    # Metadata
    created = dateutil_parser.parse(exp["created_at"])
    meta_style = ParagraphStyle("meta", parent=styles["Normal"], fontSize=10, textColor=colors.gray, spaceAfter=16)
    story.append(Paragraph(f"Created: {created.strftime('%B %d, %Y')}  ·  Status: {exp['status']}", meta_style))
    story.append(Spacer(1, 8))

    for sub in subs:
        sub_entries = [e for e in entries if e["sub_experiment_id"] == sub["id"]]

        story.append(Paragraph(sub["name"], styles["Heading2"]))
        story.append(Spacer(1, 6))

        if not sub_entries:
            story.append(Paragraph("No data recorded.", styles["Normal"]))
            story.append(Spacer(1, 12))
            continue

        headers = ["Field", "Value", "Unit", "Time", "Source"]
        if include_transcripts:
            headers.append("Transcript")

        rows = [headers]
        for e in sub_entries:
            t = dateutil_parser.parse(e["created_at"])
            row = [
                e["field_name"],
                e["field_value"],
                e["unit"] or "—",
                t.strftime("%H:%M:%S"),
                e["source"],
            ]
            if include_transcripts:
                row.append(e.get("raw_transcript") or "")
            rows.append(row)

        col_widths = [80, 60, 50, 60, 50]
        if include_transcripts:
            col_widths.append(180)

        table = Table(rows, colWidths=col_widths)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d4ed8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(table)
        story.append(Spacer(1, 20))

    doc.build(story)
    return tmp.name


def build_docx(experiment_id: str, include_transcripts: bool) -> str:
    exp, subs, entries = fetch_export_data(experiment_id, include_transcripts)

    doc = Document()

    title = doc.add_heading(exp["title"], level=0)
    title.runs[0].font.color.rgb = RGBColor(0x1d, 0x4e, 0xd8)

    created = dateutil_parser.parse(exp["created_at"])
    meta = doc.add_paragraph(f"Created: {created.strftime('%B %d, %Y')}  ·  Status: {exp['status']}")
    meta.runs[0].font.size = Pt(10)
    meta.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

    for sub in subs:
        sub_entries = [e for e in entries if e["sub_experiment_id"] == sub["id"]]
        doc.add_heading(sub["name"], level=1)

        if not sub_entries:
            doc.add_paragraph("No data recorded.")
            continue

        headers = ["Field", "Value", "Unit", "Source", "Time"]
        if include_transcripts:
            headers.append("Transcript")

        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"

        hdr = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr[i].text = h
            hdr[i].paragraphs[0].runs[0].font.bold = True

        for e in sub_entries:
            t = dateutil_parser.parse(e["created_at"])
            row = table.add_row().cells
            row[0].text = e["field_name"]
            row[1].text = e["field_value"]
            row[2].text = e["unit"] or "—"
            row[3].text = e["source"]
            row[4].text = t.strftime("%H:%M:%S")
            if include_transcripts:
                row[5].text = e.get("raw_transcript") or ""

        doc.add_paragraph()

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(tmp.name)
    return tmp.name


@router.get("/export/{experiment_id}")
def export_experiment(
    experiment_id: str,
    format: str = "pdf",
    include_transcripts: bool = False
):
    try:
        if format == "pdf":
            path = build_pdf(experiment_id, include_transcripts)
            media_type = "application/pdf"
            filename = "experiment_report.pdf"
        elif format == "docx":
            path = build_docx(experiment_id, include_transcripts)
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            filename = "experiment_report.docx"
        else:
            raise HTTPException(status_code=400, detail="format must be pdf or docx")

        return FileResponse(path, media_type=media_type, filename=filename)

    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Export error: {str(e)}")